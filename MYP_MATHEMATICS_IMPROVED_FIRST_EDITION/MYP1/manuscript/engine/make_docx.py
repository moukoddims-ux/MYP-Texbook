"""Export the manuscript to an editable Word document.

The Word file mirrors the structure of the PDF: the same chapters, sections,
worked examples, practice sets, tables and figures, with the answers collected
at the end of each chapter.
"""
import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import chaptermeta
import figs_registry
from bookkit import style as bst

bst.register_fonts()

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(os.path.dirname(ROOT), "final")
WORK = os.path.join(os.path.dirname(ROOT), "working")

GREEK = {"\\alpha": "α", "\\beta": "β", "\\pi": "π", "\\theta": "θ", "\\times": "×",
         "\\div": "÷", "\\pm": "±", "\\approx": "≈", "\\ne": "≠", "\\le": "≤",
         "\\ge": "≥", "\\sqrt": "√", "\\degree": "°", "\\circ": "°", "\\cdot": "·",
         "\\in": "∈", "\\notin": "∉", "\\cup": "∪", "\\cap": "∩", "\\emptyset": "∅",
         "\\parallel": "∥", "\\perp": "⊥", "\\angle": "∠", "\\triangle": "△",
         "\\rightarrow": "→", "\\dots": "…", "\\%": "%", "\\{": "{", "\\}": "}",
         "\\ ": " ", ",": " ", "\\square": "□"}

SUP = str.maketrans("0123456789+-n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻ⁿ")


def _math(text):
    """Convert the small maths vocabulary used in the manuscript to Unicode."""
    s = text
    for k in sorted(GREEK, key=len, reverse=True):
        s = s.replace(k, GREEK[k])
    # \frac{a}{b} -> a/b  (with brackets when either part is compound)
    def frac(m):
        a, b = m.group(1), m.group(2)
        a = a if re.fullmatch(r"[\w.]+", a) else "(%s)" % a
        b = b if re.fullmatch(r"[\w.]+", b) else "(%s)" % b
        return "%s/%s" % (a, b)
    for _ in range(3):
        s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", frac, s)
    # powers
    s = re.sub(r"\^\{?([0-9n+\-]{1,3})\}?", lambda m: m.group(1).translate(SUP), s)
    s = re.sub(r"_\{?([0-9a-z]{1,3})\}?", lambda m: m.group(1), s)
    s = re.sub(r"\\([A-Za-z]+)", lambda m: m.group(1), s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("  ", " ")
    return s.strip()


def plain(text):
    """Convert manuscript markup to plain text for Word."""
    out = []
    pos = 0
    for m in re.finditer(r"\$([^$]*)\$", text):
        out.append(_inline(text[pos:m.start()]))
        out.append(_math(m.group(1)))
        pos = m.end()
    out.append(_inline(text[pos:]))
    return "".join(out)


def _inline(chunk):
    chunk = re.sub(r"\*\*(.+?)\*\*", r"\1", chunk)
    chunk = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", chunk)
    chunk = chunk.replace("\\text{", "").replace("\\times", "×")
    chunk = re.sub(r"\\([A-Za-z]+)", r"\1", chunk)
    return chunk.replace("\\\\", "\\")


def parse_file(path):
    """Return a list of (tag, fields, body_lines)."""
    blocks = []
    cur = None
    with open(path, encoding="utf-8") as fh:
        for raw in fh:
            ln = raw.rstrip("\n")
            if ln.startswith("@") and re.match(r"^@[A-Z0-9_]+\b", ln):
                tag = ln.split()[0][1:]
                rest = ln[len(tag) + 1:].strip()
                fields = [f.strip() for f in rest.lstrip("|").split("|")] if rest else []
                cur = [tag, fields, []]
                blocks.append(cur)
            elif cur is not None:
                cur[2].append(ln)
    return [tuple(b) for b in blocks]


FIG_CACHE = {}


def figure_png(fid):
    if fid in FIG_CACHE:
        return FIG_CACHE[fid]
    entry = figs_registry.FIGS.get(fid)
    if not entry:
        return None
    from reportlab.graphics import renderPDF
    drawing = entry[0]()
    tmp_pdf = os.path.join(WORK, "fig_%s.pdf" % fid)
    renderPDF.drawToFile(drawing, tmp_pdf)
    import pymupdf
    doc = pymupdf.open(tmp_pdf)
    pix = doc[0].get_pixmap(dpi=200)
    png = os.path.join(WORK, "fig_%s.png" % fid)
    pix.save(png)
    FIG_CACHE[fid] = (png, entry[1])
    return FIG_CACHE[fid]


def add_para(doc, text, style=None, size=None, bold=False, italic=False,
             space_before=0, space_after=4):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_table(doc, rows, caption=""):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells = t.add_row().cells
        for i in range(ncol):
            txt = plain(r[i]) if i < len(r) else ""
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            run = para.add_run(txt)
            run.font.size = Pt(9)
            if r is rows[0]:
                run.bold = True
    if caption:
        add_para(doc, plain(caption), size=8.5, italic=True, space_after=8)
    else:
        doc.add_paragraph()


def chapter_doc(doc, ch, path, answers):
    doc.add_page_break()
    doc.add_heading("%d  %s" % (ch["n"], ch["title"]), level=1)
    add_para(doc, plain(ch.get("soi", "")), size=9.5, italic=True, space_after=8)
    if ch.get("inquiry"):
        add_para(doc, "Inquiry: " + plain(ch["inquiry"]), size=9, space_after=10)
    for tag, fields, body in parse_file(path):
        f = fields
        if tag == "MOVE":
            letter = f[0] if f else ""
            title = f[1] if len(f) > 1 else ""
            sub = f[2] if len(f) > 2 else ""
            doc.add_heading("%s · %s%s" % (letter, title, (" — " + sub) if sub else ""), level=2)
        elif tag == "SEC":
            num = f[0] if f else ""
            title = f[1] if len(f) > 1 else ""
            crit = f[2] if len(f) > 2 else ""
            doc.add_heading(("%s  %s" % (num, title)).strip(), level=3)
            if crit:
                add_para(doc, crit, size=8, italic=True, space_after=6)
        elif tag in ("H2", "H3", "MINI"):
            doc.add_heading(plain(f[0] if f else ""), level=4)
        elif tag == "LEAD":
            txt = " ".join(l.strip() for l in body if l.strip()) or \
                  " ".join(f for f in fields if f)
            if txt:
                add_para(doc, txt, size=10.5, italic=True, space_after=8)
        elif tag in ("P", "DIV"):
            txt = " ".join(l.strip() for l in body if l.strip()) or \
                  " ".join(f for f in fields if f)
            if txt:
                add_para(doc, txt, size=10, space_after=6)
        elif tag in ("DEF", "KEY", "ERR", "ATL", "NOTE", "METHOD", "INVEST", "APPLY",
                     "CRIT", "REFLECT", "CALLOUT"):
            title = f[0] if f else ""
            label = {"DEF": "DEFINITION", "KEY": "KEY IDEA", "ERR": "COMMON ERROR",
                     "ATL": "APPROACHES TO LEARNING", "NOTE": "NOTE", "METHOD": "METHOD",
                     "INVEST": "INVESTIGATION", "APPLY": "APPLICATION",
                     "CRIT": "CRITERION TASK", "CALLOUT": "INQUIRY",
                     "REFLECT": "REFLECTION"}.get(tag, tag)
            add_para(doc, "%s · %s" % (label, plain(title)), size=8, bold=True, space_before=8)
            for l in body:
                q = l.strip()
                if not q:
                    continue
                q = re.sub(r"^\d+\.\s*", "", q)
                q = re.sub(r"^-\s*", "", q)
                add_para(doc, plain(q), size=9.5, space_after=3)
        elif tag == "EX":
            ref = f[0] if f else ""
            title = f[1] if len(f) > 1 else ""
            add_para(doc, "WORKED EXAMPLE %s · %s" % (ref, plain(title)), size=10, bold=True,
                     space_before=10)
            for l in body:
                q = l.strip()
                if q:
                    add_para(doc, plain(q), size=9.5, space_after=3)
        elif tag == "PRACTICE":
            pid = f[0] if f else ""
            ptitle = f[1] if len(f) > 1 else ""
            levels = f[2] if len(f) > 2 else ""
            add_para(doc, "PRACTICE %s · %s   (%s)" % (pid, plain(ptitle), plain(levels)),
                     size=9.5, bold=True, space_before=10)
            num = 0
            items = []
            for l in body:
                q = l.strip()
                if not q.startswith("- "):
                    continue
                parts = q[2:].split("||")
                left = parts[0]
                ans = parts[1].strip() if len(parts) > 1 else ""
                bits = [b.strip() for b in left.split("|")]
                typ = bits[0]
                text = " | ".join(bits[1:]) if len(bits) > 1 else left
                num += 1
                add_para(doc, "%d.  [%s]  %s" % (num, typ, plain(text)), size=9.5, space_after=2)
                items.append((num, ans))
            answers.append((pid, ptitle, items))
        elif tag == "TABLE":
            caption = f[0] if f else ""
            rows = []
            for l in body:
                q = l.strip()
                if q and not q.startswith("@"):
                    rows.append([c.strip() for c in q.split(";")])
            add_table(doc, rows, caption)
        elif tag == "FIG":
            fid = f[0] if f else ""
            got = figure_png(fid)
            if got:
                png, cap = got
                doc.add_picture(png, width=Inches(4.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_para(doc, "Figure: " + plain(cap), size=8.5, italic=True, space_after=8)
        elif tag == "LISTU" or tag == "LISTO":
            n = 0
            for l in body:
                q = l.strip()
                if not q:
                    continue
                q = re.sub(r"^\d+\.\s*", "", q)
                q = re.sub(r"^-\s*", "", q)
                if tag == "LISTO":
                    n += 1
                    add_para(doc, "%d.  %s" % (n, plain(q)), size=9.5, space_after=2)
                else:
                    add_para(doc, "•  %s" % plain(q), size=9.5, space_after=2)
        elif tag == "EQ":
            eqs = [plain("$%s$" % x) if not x.strip().startswith("$") else plain(x)
                   for x in (f + [l.strip() for l in body if l.strip()]) if x.strip()]
            for e in eqs:
                p = add_para(doc, e, size=10.5, space_after=6)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tag in ("SPACE", "PAGEBREAK", "IDX", "RAWSPACER", "H1"):
            continue
        else:
            continue
        # stray body text for inline tags
    if answers:
        doc.add_heading("Answers to Chapter %d" % ch["n"], level=3)
        for pid, ptitle, items in answers:
            add_para(doc, "Practice %s · %s" % (pid, plain(ptitle)), size=9, bold=True,
                     space_before=6)
            for n, a in items:
                if a:
                    add_para(doc, "%d.  %s" % (n, plain(a)), size=8.8, space_after=1)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10)
    sec = doc.sections[0]
    from docx.shared import Cm
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MYP MATHEMATICS 1")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = RGBColor(0x14, 0x2B, 0x3A)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Student Coursebook · Improved First Edition")
    r.font.size = Pt(15)
    r.italic = True
    auth = doc.add_paragraph()
    auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = auth.add_run("SADIKOU MOUKODDIM")
    r.font.size = Pt(13)
    r.bold = True
    add_para(doc, "Editable Word edition, exported from the same manuscript as the PDF. "
                  "This is an improved first edition, not a second edition. It is an independent "
                  "publication and is not published, sponsored or endorsed by the International "
                  "Baccalaureate Organization. Copyright © SADIKOU Moukoddim. All rights reserved.",
             size=9, space_before=18)

    doc.add_page_break()
    doc.add_heading("How to use this book", level=1)
    for line in [
        "Each chapter follows the same rhythm: A Engage, B Explore, C Explain, "
        "D Worked examples, E Practice, F Investigate and apply, G Reflect and consolidate.",
        "Definitions, key ideas and common errors appear in coloured boxes. Each box is labelled, "
        "so you can find the rule again when you need it.",
        "Practice sets are differentiated: SUPPORT questions build the method, CORE questions "
        "make it dependable, and EXTENSION questions use it in an unfamiliar setting.",
        "An investigation (criterion B) and an application (criterion D) close each chapter, "
        "with success criteria printed beside them. The Task C writing task develops "
        "communication (criterion C).",
        "Answers are collected at the back of the PDF edition and at the end of each chapter in "
        "this Word edition, so that you can check your reasoning rather than replace it.",
    ]:
        add_para(doc, line, size=10, space_after=5)

    doc.add_heading("Assessment criteria A–D", level=1)
    crit = [
        ("A  Knowing and understanding",
         "Select and apply mathematical knowledge, techniques and notation accurately."),
        ("B  Investigating patterns",
         "Choose and apply techniques to explore a pattern, describe it generally and verify it."),
        ("C  Communicating",
         "Use appropriate mathematical language, notation and representation to explain reasoning."),
        ("D  Applying mathematics in real-life contexts",
         "Identify relevant elements, choose a model, apply it, and justify and evaluate the result."),
    ]
    for name, text in crit:
        add_para(doc, name, size=10.5, bold=True, space_before=6, space_after=2)
        add_para(doc, text, size=10, space_after=4)
    add_para(doc, "The level descriptors used beside tasks in this book are teaching tools written "
                  "for this course. They paraphrase the MYP criteria and do not replace the current "
                  "MYP Mathematics guide.", size=9, italic=True, space_before=6)

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    for ch in chaptermeta.CHAPTERS:
        add_para(doc, "%d   %s" % (ch["n"], ch["title"]), size=10, space_after=2)

    for ch in chaptermeta.CHAPTERS:
        path = os.path.join(ROOT, ch["content"])
        answers = []
        chapter_doc(doc, ch, path, answers)

    # cumulative revision
    rev = os.path.join(ROOT, "content", "revision.txt")
    if os.path.exists(rev):
        doc.add_page_break()
        doc.add_heading("Cumulative revision and part reviews", level=1)
        chapter_doc(doc, {"n": 0, "title": "Cumulative revision", "soi": ""}, rev, [])

    doc.add_page_break()
    doc.add_heading("Glossary", level=1)
    import chapterdata
    for term, meaning in chapterdata.GLOSSARY:
        add_para(doc, "%s — %s" % (term, plain(meaning)), size=9, space_after=2)

    os.makedirs(OUT, exist_ok=True)
    out = os.path.join(OUT, "MYP1_Mathematics_Improved_First_Edition.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
